from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
EXCEL_DIR = ROOT / "Excel"
JSON_PATH = ROOT / "data" / "tables" / "cards.json"

ATTR_SLOT_COUNT = 3
SKILL_SLOT_COUNT = 3
ATTR_LINE_RE = re.compile(r"^(.+?)\s*([+-])\s*(\d+(?:\.\d+)?)\s*(%)?$")

HEADERS = [
    "id",
    "name",
    "icon",
    "tier",
    "is_unique",
    "description",
    "attr_1_key",
    "attr_1_op",
    "attr_1_value",
    "attr_1_is_percent",
    "attr_1_text",
    "attr_2_key",
    "attr_2_op",
    "attr_2_value",
    "attr_2_is_percent",
    "attr_2_text",
    "attr_3_key",
    "attr_3_op",
    "attr_3_value",
    "attr_3_is_percent",
    "attr_3_text",
    "skill_1_id",
    "skill_1_text",
    "skill_2_id",
    "skill_2_text",
    "skill_3_id",
    "skill_3_text",
    "stat",
    "value",
    "notes",
]

TYPE_ROW = [
    "string",
    "string",
    "string",
    "int",
    "bool",
    "string",
    "string",
    "string",
    "string",
    "bool",
    "string",
    "string",
    "string",
    "string",
    "bool",
    "string",
    "string",
    "string",
    "string",
    "bool",
    "string",
    "string",
    "string",
    "string",
    "string",
    "string",
    "string",
    "string",
    "string",
    "string",
]

LABEL_ROW = [
    "ID",
    "名称",
    "图标",
    "阶级",
    "是否唯一",
    "描述",
    "属性1-键",
    "属性1-运算",
    "属性1-数值",
    "属性1-百分比",
    "属性1-文本",
    "属性2-键",
    "属性2-运算",
    "属性2-数值",
    "属性2-百分比",
    "属性2-文本",
    "属性3-键",
    "属性3-运算",
    "属性3-数值",
    "属性3-百分比",
    "属性3-文本",
    "技能1-ID(占位)",
    "技能1-文本",
    "技能2-ID(占位)",
    "技能2-文本",
    "技能3-ID(占位)",
    "技能3-文本",
    "兼容旧字段(stat)",
    "兼容旧字段(value)",
    "备注/待整理项",
]


def find_workbook_path() -> Path:
    matches = sorted(EXCEL_DIR.glob("cards*.xlsx"))
    if not matches:
        raise FileNotFoundError("未找到 Excel/cards*.xlsx")
    return matches[0]


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return text.strip()


def read_source_rows(path: Path) -> list[dict]:
    workbook = load_workbook(path)
    sheet = workbook[workbook.sheetnames[0]]
    header_map = {}
    for column in range(1, sheet.max_column + 1):
        header = normalize_text(sheet.cell(1, column).value)
        if header:
            header_map[header] = column

    rows: list[dict] = []
    for row_index in range(4, sheet.max_row + 1):
        card_id = normalize_text(sheet.cell(row_index, header_map["id"]).value)
        if not card_id:
            continue
        row = {
            "id": card_id,
            "name": normalize_text(sheet.cell(row_index, header_map["name"]).value),
            "icon": normalize_text(sheet.cell(row_index, header_map.get("icon", 0)).value),
            "tier": int(sheet.cell(row_index, header_map["tier"]).value or 0),
            "is_unique": bool(sheet.cell(row_index, header_map["is_unique"]).value),
            "description": normalize_text(sheet.cell(row_index, header_map["description"]).value),
            "stat": normalize_text(sheet.cell(row_index, header_map.get("stat", 0)).value),
            "value": normalize_text(sheet.cell(row_index, header_map.get("value", 0)).value),
        }
        rows.append(row)
    workbook.close()
    return rows


def is_attr_line(line: str) -> bool:
    if not line:
        return False
    if ":" in line or "：" in line:
        return False
    if line.startswith("唯一被动") or line.startswith("被动"):
        return False
    return ATTR_LINE_RE.match(line) is not None


def build_attr_entry(line: str) -> dict:
    match = ATTR_LINE_RE.match(line)
    if match is None:
        raise ValueError(f"无法解析属性行: {line}")
    return {
        "key": match.group(1).strip(),
        "op": match.group(2),
        "value": match.group(3),
        "is_percent": bool(match.group(4)),
        "text": line,
    }


def split_description(description: str) -> tuple[list[dict], list[str]]:
    attr_entries: list[dict] = []
    skill_lines: list[str] = []
    for raw_line in description.split("\n"):
        line = normalize_text(raw_line)
        if not line:
            continue
        if is_attr_line(line):
            attr_entries.append(build_attr_entry(line))
        else:
            skill_lines.append(line)
    return attr_entries[:ATTR_SLOT_COUNT], skill_lines[:SKILL_SLOT_COUNT]


def build_row(base_row: dict) -> dict:
    attr_entries, skill_lines = split_description(base_row["description"])
    new_row = {header: "" for header in HEADERS}
    new_row["id"] = base_row["id"]
    new_row["name"] = base_row["name"]
    new_row["icon"] = base_row["icon"]
    new_row["tier"] = base_row["tier"]
    new_row["is_unique"] = base_row["is_unique"]
    new_row["description"] = base_row["description"]
    new_row["stat"] = base_row["stat"]
    new_row["value"] = base_row["value"]
    new_row["notes"] = ""

    for index in range(ATTR_SLOT_COUNT):
        slot = index + 1
        entry = attr_entries[index] if index < len(attr_entries) else None
        new_row[f"attr_{slot}_key"] = entry["key"] if entry else ""
        new_row[f"attr_{slot}_op"] = entry["op"] if entry else ""
        new_row[f"attr_{slot}_value"] = entry["value"] if entry else ""
        new_row[f"attr_{slot}_is_percent"] = entry["is_percent"] if entry else False
        new_row[f"attr_{slot}_text"] = entry["text"] if entry else ""

    for index in range(SKILL_SLOT_COUNT):
        slot = index + 1
        line = skill_lines[index] if index < len(skill_lines) else ""
        new_row[f"skill_{slot}_id"] = str(slot) if line else ""
        new_row[f"skill_{slot}_text"] = line

    return new_row


def write_workbook(path: Path, rows: list[dict]) -> None:
    workbook = load_workbook(path)
    sheet = workbook[workbook.sheetnames[0]]

    if sheet.max_row:
        sheet.delete_rows(1, sheet.max_row)
    if sheet.max_column > len(HEADERS):
        sheet.delete_cols(len(HEADERS) + 1, sheet.max_column - len(HEADERS))

    sheet.append(HEADERS)
    sheet.append(TYPE_ROW)
    sheet.append(LABEL_ROW)
    for row in rows:
        sheet.append([row.get(header, "") for header in HEADERS])

    workbook.save(path)
    workbook.close()


def write_json(path: Path, rows: list[dict]) -> None:
    payload = json.dumps(rows, ensure_ascii=False, indent=4)
    path.write_text(payload + "\n", encoding="utf-8")


def main() -> None:
    workbook_path = find_workbook_path()
    source_rows = read_source_rows(workbook_path)
    rebuilt_rows = [build_row(row) for row in source_rows]
    write_workbook(workbook_path, rebuilt_rows)
    write_json(JSON_PATH, rebuilt_rows)
    print(f"重建完成: {workbook_path}")
    print(f"同步完成: {JSON_PATH}")
    print(f"卡牌数量: {len(rebuilt_rows)}")


if __name__ == "__main__":
    main()

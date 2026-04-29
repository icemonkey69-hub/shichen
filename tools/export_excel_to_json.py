from __future__ import annotations

import argparse
import json
import time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
EXCEL_DIR = ROOT / "Excel"
OUTPUT_DIR = ROOT / "data" / "tables"

HEADER_ROW = 1
TYPE_ROW = 2
DATA_START_ROW = 4

TRUE_SET = {"1", "true", "yes", "y", "on", "是", "启用"}
FALSE_SET = {"0", "false", "no", "n", "off", "否", "禁用", ""}


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export Excel tables to JSON.")
    parser.add_argument(
        "--force",
        action="store_true",
        help="Force export all tables, ignore unchanged-file skip.",
    )
    return parser.parse_args()


def _log(message: str) -> None:
    print(message, flush=True)


def _table_name_from_excel(path: Path) -> str:
    stem = path.stem
    if "#" in stem:
        stem = stem.split("#", 1)[0]
    return stem.strip()


def _normalize_type(raw: Any) -> str:
    if raw is None:
        return "string"
    t = str(raw).strip().lower()
    return t if t in {"string", "int", "float", "bool"} else "string"


def _to_string(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _to_int(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return 1 if value else 0
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)
    text = str(value).strip()
    if text == "":
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def _to_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return 1.0 if value else 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if text == "":
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _to_bool(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return float(value) != 0.0
    text = str(value).strip().lower()
    if text in TRUE_SET:
        return True
    if text in FALSE_SET:
        return False
    return False


def _convert(value: Any, type_name: str) -> Any:
    if type_name == "int":
        return _to_int(value)
    if type_name == "float":
        return _to_float(value)
    if type_name == "bool":
        return _to_bool(value)
    return _to_string(value)


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip() == ""
    return False


def _read_sheet(xlsx_path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = wb.active
    rows_preview = list(ws.iter_rows(min_row=HEADER_ROW, max_row=TYPE_ROW, values_only=True))
    header_values = rows_preview[0] if len(rows_preview) >= 1 else tuple()
    type_values = rows_preview[1] if len(rows_preview) >= 2 else tuple()

    max_col = max(len(header_values), len(type_values))
    columns: list[tuple[int, str, str]] = []
    for idx in range(max_col):
        key_raw = header_values[idx] if idx < len(header_values) else None
        key = str(key_raw).strip() if key_raw is not None else ""
        if key == "":
            continue
        type_raw = type_values[idx] if idx < len(type_values) else None
        type_name = _normalize_type(type_raw)
        columns.append((idx, key, type_name))

    rows: list[dict[str, Any]] = []
    if not columns:
        wb.close()
        return rows

    max_col_for_iter = max(idx for (idx, _, _) in columns) + 1
    for row_values in ws.iter_rows(min_row=DATA_START_ROW, max_col=max_col_for_iter, values_only=True):
        raw_cells = [row_values[idx] if idx < len(row_values) else None for (idx, _, _) in columns]
        if all(_is_blank(v) for v in raw_cells):
            continue

        item: dict[str, Any] = {}
        for (idx, key, type_name) in columns:
            raw_value = row_values[idx] if idx < len(row_values) else None
            item[key] = _convert(raw_value, type_name)
        rows.append(item)

    wb.close()
    return rows


def main() -> int:
    args = _parse_args()
    if not EXCEL_DIR.exists():
        _log(f"[ERROR] Excel directory not found: {EXCEL_DIR}")
        return 1
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    excel_files = sorted(
        p
        for p in EXCEL_DIR.glob("*.xlsx")
        if not p.name.startswith("~$")
    )
    if not excel_files:
        _log(f"[WARN] No xlsx found in: {EXCEL_DIR}")
        return 0

    exported = 0
    skipped = 0
    total = len(excel_files)
    for index, xlsx_path in enumerate(excel_files, start=1):
        table_name = _table_name_from_excel(xlsx_path)
        if table_name == "":
            _log(f"[SKIP] ({index}/{total}) invalid table name: {xlsx_path.name}")
            continue
        out_path = OUTPUT_DIR / f"{table_name}.json"
        if (
            not args.force
            and out_path.exists()
            and out_path.stat().st_mtime >= xlsx_path.stat().st_mtime
        ):
            skipped += 1
            _log(f"[SKIP] ({index}/{total}) unchanged: {xlsx_path.name}")
            continue

        _log(f"[RUN ] ({index}/{total}) {xlsx_path.name}")
        started_at = time.perf_counter()
        rows = _read_sheet(xlsx_path)
        payload = json.dumps(rows, ensure_ascii=False, indent=4)
        out_path.write_text(payload + "\n", encoding="utf-8")
        exported += 1
        elapsed = time.perf_counter() - started_at
        _log(
            f"[OK] {xlsx_path.name} -> data/tables/{table_name}.json "
            f"({len(rows)} rows, {elapsed:.2f}s)"
        )

    _log(f"[DONE] Exported {exported} table(s), skipped {skipped} unchanged table(s).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

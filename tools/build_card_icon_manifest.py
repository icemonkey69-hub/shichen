from __future__ import annotations

import argparse
import json
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
MASTER_XLSX = next(ROOT.glob("*.xlsx"))
CARDS_JSON = ROOT / "data" / "tables" / "cards.json"
OUT_JSON = ROOT / "tools" / "card_icon_manifest.json"


def load_master_prompt_map() -> dict[str, dict]:
    workbook = openpyxl.load_workbook(MASTER_XLSX, read_only=True)
    worksheet = workbook[workbook.sheetnames[4]]
    prompt_map: dict[str, dict] = {}
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        name, tier, description, prompt = row[:4]
        if not name:
            continue
        prompt_map[str(name).strip()] = {
            "tier": int(tier or 0),
            "description": str(description or "").strip(),
            "prompt": str(prompt or "").strip(),
        }
    return prompt_map


def tier_background_prompt(tier: int) -> str:
    mapping = {
        1: "Background color should stay close to tier 1: white, silver white, pale gray white, cool misty white",
        2: "Background color should stay close to tier 2: green, emerald green, forest green, teal green",
        3: "Background color should stay close to tier 3: blue, deep blue, icy blue, azure blue",
        4: "Background color should stay close to tier 4: purple, arcane purple, dark violet, mystical violet",
        5: "Background color should stay close to tier 5: orange, molten gold orange, amber orange, flame orange",
        6: "Background color should stay close to tier 6: red, crimson, blood red, apocalyptic dark red",
    }
    return mapping.get(tier, "Background color should stay aligned with the card tier and maintain a clean unified fantasy atmosphere")


def build_single_icon_prompt(name: str, prompt: str, tier: int) -> str:
    base_prompt = prompt.strip() or name.strip()
    background_prompt = tier_background_prompt(tier)
    return (
        "Game card icon generation. "
        "Create a single 512x512 square low-poly divine-demonic fantasy game icon. "
        "The subject must stay centered, fully visible, with safe padding on all sides. "
        "Do not place the subject against the edge, do not crop it, and do not push it into a corner. "
        "Keep the image clean and readable at small size. "
        "Visual style: polished low-poly 3D, mythic divine-demonic world, clear chunky forms, one clear focal subject, restrained glow, clean background. "
        "No text, no borders, no UI decorations, no realistic photography, no multiple unrelated subjects, no complex scene. "
        f"{background_prompt}. "
        f"Main subject: {base_prompt}"
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--tier", type=int, action="append", help="only export matching card tiers; can be repeated")
    parser.add_argument("--icon", action="append", help="only export matching icon ids; can be repeated")
    parser.add_argument("--limit", type=int, default=0, help="max number of rows to export after filtering")
    parser.add_argument("--out", default=str(OUT_JSON), help="output json path")
    args = parser.parse_args()

    cards = json.loads(CARDS_JSON.read_text(encoding="utf-8"))
    prompt_map = load_master_prompt_map()
    manifest: list[dict] = []
    allowed_tiers = set(args.tier or [])
    allowed_icons = {str(icon_id).strip() for icon_id in (args.icon or []) if str(icon_id).strip()}

    for index, card in enumerate(cards, start=1):
        name = str(card.get("name", "")).strip()
        icon_id = str(card.get("icon", "")).strip()
        if not icon_id:
            continue
        prompt_info = prompt_map.get(name, {})
        prompt = str(prompt_info.get("prompt", "")).strip()
        description = str(prompt_info.get("description", "")).strip() or str(card.get("description", "")).strip()
        tier = int(prompt_info.get("tier", card.get("tier", 0)) or 0)
        if allowed_tiers and tier not in allowed_tiers:
            continue
        if allowed_icons and icon_id not in allowed_icons:
            continue
        manifest.append(
            {
                "index": index,
                "id": str(card.get("id", "")).strip(),
                "icon": icon_id,
                "name": name,
                "tier": tier,
                "description": description,
                "raw_prompt": prompt,
                "single_icon_prompt": build_single_icon_prompt(name, prompt or description or name, tier),
                "target_path": str(ROOT / "assets" / "ui" / "icons" / "cards" / f"{icon_id}.png"),
            }
        )
        if args.limit > 0 and len(manifest) >= args.limit:
            break

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"wrote {len(manifest)} icon prompt rows -> {out_path}")


if __name__ == "__main__":
    main()

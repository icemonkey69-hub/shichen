from __future__ import annotations

import json
import math
import openpyxl
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
MASTER_XLSX = next(ROOT.glob("*.xlsx"))
CARDS_JSON = ROOT / "data" / "tables" / "cards.json"
OUT_JSON = ROOT / "tools" / "card_icon_batches.json"

BATCH_SIZE = 6


def load_master_prompt_map() -> dict[str, dict]:
    wb = openpyxl.load_workbook(MASTER_XLSX, read_only=True)
    ws = wb[wb.sheetnames[4]]
    prompt_map: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        name, tier, description, prompt = row[:4]
        if not name:
            continue
        prompt_map[str(name).strip()] = {
            "tier": int(tier or 0),
            "description": str(description or "").strip(),
            "prompt": str(prompt or "").strip(),
        }
    return prompt_map


def main() -> None:
    cards = json.loads(CARDS_JSON.read_text(encoding="utf-8"))
    prompt_map = load_master_prompt_map()
    rows: list[dict] = []
    for card in cards:
        name = str(card.get("name", "")).strip()
        icon = str(card.get("icon", "")).strip()
        info = prompt_map.get(name, {})
        rows.append(
            {
                "id": str(card.get("id", "")),
                "name": name,
                "icon": icon,
                "tier": int(card.get("tier", 0)),
                "description": str(card.get("description", "")).strip(),
                "prompt": str(info.get("prompt", "")).strip(),
            }
        )

    batches: list[dict] = []
    total_batches = int(math.ceil(len(rows) / float(BATCH_SIZE)))
    for batch_index in range(total_batches):
        items = rows[batch_index * BATCH_SIZE : (batch_index + 1) * BATCH_SIZE]
        layout_lines = []
        for idx, item in enumerate(items, start=1):
            prompt = item["prompt"] or item["description"] or item["name"]
            layout_lines.append(f"{idx}. {item['icon']} | {item['name']} | {prompt}")
        batches.append(
            {
                "batch_index": batch_index + 1,
                "count": len(items),
                "items": items,
                "layout_lines": layout_lines,
            }
        )

    OUT_JSON.write_text(json.dumps(batches, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"wrote {len(batches)} batches -> {OUT_JSON}")


if __name__ == "__main__":
    main()
